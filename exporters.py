"""Convert the Markdown report into Word (.docx) and PDF files (as bytes)."""
import io
import re
from pathlib import Path

import markdown as md_lib
from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF

# ----------------------------------------------------------------------
# Helpers shared by both exporters
# ----------------------------------------------------------------------
LIST_LINE = re.compile(r"^\s*([-*+]|\d+[.)])\s+")
INLINE = re.compile(
    r"\*\*(?P<b>.+?)\*\*|`(?P<c>.+?)`|\*(?P<i>[^*\s][^*]*?)\*"
    r"|\[(?P<lt>.+?)\]\((?P<lu>.+?)\)"
)


def _normalize_md(text: str) -> str:
    """LLMs often put a list right after a text line without a blank line.
    The Markdown library needs that blank line, so we add it."""
    out, in_code = [], False
    prev = ""
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        if (not in_code and LIST_LINE.match(line) and prev.strip()
                and not LIST_LINE.match(prev)):
            out.append("")
        out.append(line)
        prev = line
    return "\n".join(out)


def _strip_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", text)
    return text


# ----------------------------------------------------------------------
# WORD (.docx)
# ----------------------------------------------------------------------
def _add_inline(paragraph, text: str):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group("b") is not None:
            paragraph.add_run(m.group("b").replace("`", "")).bold = True
        elif m.group("c") is not None:
            run = paragraph.add_run(m.group("c"))
            run.font.name = "Consolas"
        elif m.group("i") is not None:
            paragraph.add_run(m.group("i")).italic = True
        else:
            paragraph.add_run(f"{m.group('lt')} ({m.group('lu')})")
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _bullet_level(leading_spaces: int) -> int:
    return 0 if leading_spaces < 2 else (1 if leading_spaces < 6 else 2)


def markdown_to_docx(md_text: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue

        # code block
        if s.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                run = p.add_run(lines[i])
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                i += 1
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            doc.add_heading(_strip_inline(m.group(2)), level=min(len(m.group(1)), 4))
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"([-*_]\s*){3,}", s):
            i += 1
            continue

        # table
        if s.startswith("|") and i + 1 < len(lines) and re.match(
                r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(cols):
                    cell = table.cell(r, c)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    _add_inline(para, row[c] if c < len(row) else "")
                    if r == 0:
                        for run in para.runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        # bullet list item
        m = re.match(r"^(\s*)[-*+]\s+(.*)", line)
        if m:
            level = _bullet_level(len(m.group(1).replace("\t", "    ")))
            style = "List Bullet" if level == 0 else f"List Bullet {level + 1}"
            _add_inline(doc.add_paragraph(style=style), m.group(2))
            i += 1
            continue

        # numbered list item (kept as plain text so numbering never "continues")
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)", line)
        if m:
            level = _bullet_level(len(m.group(1).replace("\t", "    ")))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3 + 0.3 * level)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            _add_inline(p, f"{m.group(2)}. {m.group(3)}")
            i += 1
            continue

        # normal paragraph (join wrapped lines)
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns or ns.startswith(("#", "```", "|")) or LIST_LINE.match(nxt)):
                break
            buf.append(ns)
            i += 1
        _add_inline(doc.add_paragraph(), " ".join(buf))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
FONT_DIRS = [
    Path("/usr/share/fonts/truetype/dejavu"),
    Path("/usr/share/fonts/dejavu"),
    Path("C:/Windows/Fonts"),
]

CHAR_FIXES = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-",
    "\u2192": "->", "\u2190": "<-", "\u20b9": "Rs.", "\u00a0": " ",
}


def _find_font(*names: str):
    for folder in FONT_DIRS:
        for name in names:
            path = folder / name
            if path.exists():
                return str(path)
    return None


def _latin1_safe(text: str) -> str:
    for old, new in CHAR_FIXES.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


def markdown_to_pdf(md_text: str) -> bytes:
    html = md_lib.markdown(
        _normalize_md(md_text),
        extensions=["tables", "fenced_code", "sane_lists"],
    )

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()

    regular = _find_font("DejaVuSans.ttf", "dejavusans.ttf")
    bold = _find_font("DejaVuSans-Bold.ttf", "dejavusans-bold.ttf") or regular
    mono = _find_font("DejaVuSansMono.ttf", "dejavusansmono.ttf") or regular

    if regular:   # full Unicode support
        pdf.add_font("DejaVu", "", regular)
        pdf.add_font("DejaVu", "B", bold)
        pdf.add_font("DejaVu", "I", regular)     # no italic file: reuse regular
        pdf.add_font("DejaVu", "BI", bold)
        pdf.add_font("DejaVuMono", "", mono)
        pdf.add_font("DejaVuMono", "B", mono)
        family, code_font = "DejaVu", "DejaVuMono"
    else:         # fallback: built-in fonts (Latin-1 characters only)
        html = _latin1_safe(html)
        family, code_font = "helvetica", "courier"

    pdf.set_font(family, size=11)
    pdf.write_html(
        html,
        font_family=family,
        pre_code_font=code_font,
        ul_bullet_char="-",
        li_prefix_color=(60, 60, 60),
        table_line_separators=True,
        warn_on_tags_not_matching=False,
    )
    return bytes(pdf.output())
